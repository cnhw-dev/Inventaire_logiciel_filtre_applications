"""
appliquer_regle06.py — Applique la Règle 6 (exclusions spécifiques de
composantes Autodesk, Adobe, Citrix, Python et autres) sur les onglets de
secteurs de « User_Workstation_App_Associations_Consolidated.xlsx »
(jusqu'à l'onglet « Non défini » inclusivement).

Règle (conforme à Regle_06_exclusions_specifiques.md) :
les applications listées sont des composantes installées avec l'application
principale (services, modules linguistiques, greffins, composantes
d'installation Python, etc.) et sont exclues de l'inventaire. Chaque ligne
correspondante est déplacée vers l'onglet « Excluded Regle06 » sans
transformation (aucune perte de données — une colonne finale « Raison »
est ajoutée dans l'onglet d'exclusion).

Correspondance :
- la présence d'une étoile (*) dans un motif correspond à un « wildcard »
  (n'importe quelle chaîne de caractères, y compris vide);
- sans étoile, la correspondance est exacte (après normalisation) : ainsi
  « Autodesk App Manager » n'exclut que ce libellé exact, et les motifs
  Citrix exigeant la parenthèse (« Citrix Receiver (* ») n'excluent pas
  « Citrix Receiver 4.12 » ni « Citrix Workspace 2503/2507/2511 »;
- normalisation : casse ignorée, apostrophes unifiées, espaces superflus
  réduits (« autodesk  access » = « Autodesk Access »);
- les motifs Python se terminent par « * » afin de couvrir le suffixe
  « (64-bit) » des libellés réels (ex. « Python 3.11.0 Core Interpreter
  (64-bit) »).

Ajouts à la liste du MD (validés avec l'utilisateur) :
- « AutoCAD LT 2026 - Français (French) », « Autodesk AutoCAD LT 2026 -
  Français (French) » et « Autodesk CER » (équivalents LT/rapport d'erreurs
  des composantes AutoCAD 2022 exclues);
- « Citrix Browser Content Redirection Plug-in » (= « BCR Plug-in »);
- « Python * Add to Path* » et « Python * Utility Scripts* »
  (« Python Launcher » est CONSERVÉ, décision de portée);
- « Microsoft Windows Application Compatibility Fix Database »
  (composante système Microsoft).

Retrait de la liste du MD (validé avec l'utilisateur) :
- « Autodesk Access » est CONSERVÉ : bien que listé dans le MD, ce
  composant (client de mise à jour Autodesk) ne doit pas être exclu.

Comportement (calqué sur appliquer_regle05.py) :
1. Génère d'abord la liste détaillée des applications correspondantes
   (aucun déplacement), en lecture seule (read_only) pour être rapide et
   économe en mémoire. Ce rapport permet de valider les correspondances
   (et les motifs sans correspondance) avant l'exécution réelle.
2. Traite ensuite UN secteur à la fois : le classeur est rouvert pour cet
   onglet seulement, les lignes correspondantes sont déplacées vers
   l'onglet « Excluded Regle06 » (créé une seule fois), puis le classeur
   est sauvegardé, fermé et les ressources sont libérées avant de passer
   au secteur suivant.
3. La progression est affichée à l'écran (onglet courant, temps écoulé et
   estimation du temps restant). Aucune intervention n'est requise entre
   les secteurs.

Usage : py appliquer_regle06.py --dry-run   # simulation (aucune écriture)
        py appliquer_regle06.py             # exécution réelle
"""

import argparse
import gc
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from colonnes import aligner, index_colonnes, valeur

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
EXCLUDED_SHEET = "Excluded Regle06"
REASON = ("Règle 6 — Exclusion spécifique : composante installée avec "
          "l'application principale")
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
IGNORED = {"logiciels", "sommaire", "feuil1", "excluded regle01", "excluded regle02",
           "excluded regle03", "excluded regle04", "excluded regle05",
           EXCLUDED_SHEET.casefold()}


def _norm(value):
    """Normalise une valeur : casefold, apostrophes unifiées, espaces réduits."""
    text = "" if value is None else str(value)
    for ch in ("’", "‘", "`"):
        text = text.replace(ch, "'")
    return " ".join(text.split()).casefold()


def _compiler_motif(motif):
    """Convertit un motif en regex compilée (« * » = n'importe quelle chaîne).

    La correspondance est ancrée (chaîne complète) : sans étoile, le motif
    doit correspondre exactement au libellé normalisé.
    """
    expr = "".join(".*" if ch == "*" else re.escape(ch)
                   for ch in _norm(motif))
    return re.compile(f"^{expr}$")


# Motifs d'exclusion (Règle 6) par catégorie. L'étoile (*) est un
# « wildcard » : n'importe quelle chaîne de caractères (y compris vide).
# Liste dédoublonnée (« Adobe Genuine Service », « Citrix Authentication
# Manager » et « Citrix Web Helper » apparaissent en double dans le MD).
_MOTIFS = {
    "Autodesk": [
        "ACA & MEP 2022 Object Enabler",
        "Adobe Genuine Service",  # aussi listé sous « Adobe » dans le MD
        "Applications recommandées Autodesk",
        "AutoCAD 2022 - Français (French)",
        "AutoCAD 2022 Language Pack - Français (French)",
        "AutoCAD LT 2026 - Français (French)",          # ajout validé
        "AutoCAD Open in Desktop",
        # « Autodesk Access » est CONSERVÉ (retiré de la liste du MD,
        # décision de portée).
        "Autodesk App Manager",
        "Autodesk AutoCAD 2022 - Français (French)",
        "Autodesk AutoCAD LT 2026 - Français (French)",  # ajout validé
        "Autodesk CER",                                 # ajout validé
        "Autodesk Genuine Service",
        "Autodesk Identity Manager",
        "Autodesk Material Library 2022",
        "Autodesk Material Library Base Resolution Image Library 2022",
        "Autodesk Single Sign On Component",
        "Enregistrement Autodesk sur le Web et au format mobile",
    ],
    "Adobe": [
        "Adobe Native Client",
        "COCM_1_0_32",
        "COMP_1_4_32",
        "COPS_1_0_32",
        "CORE_1_0_32",
        "CORG_1_1_32",
        "COSY_7_8_10_32",
        "CCXP*",
        "KANC*",
        "KANE_1_2_0",
        "KASU_2_0_32",
        "LIBS*",
        "SEPS_0_25_1",
        "UAM_1_50",
        "*Notification Client",
    ],
    "Citrix": [
        "AppDataProtection",
        "AppProtection",
        "BCR Plug-in",
        "Citrix Authentication Manager",
        "Citrix Browser Content Redirection Plug-in",  # ajout validé (= BCR)
        "Citrix Desktop Lock",
        "Citrix Receiver (*",
        "Citrix Receiver(*",
        "Citrix Receiver Inside",
        "Citrix Secure Access Endpoint Analysis",
        "Citrix Web Helper",
        "Citrix Workspace (*",
        "Citrix Workspace(*",
        "Citrix Workspace Inside",
        "deviceTRUST ICA Client Extension*",
        "MTOP Client",
        "Online Plug-in",
        "Self-Service Plug-in",
    ],
    "Python": [
        "Python * Add to Path*",            # ajout validé
        "Python * Core Interpreter*",
        "Python * Development Libraries*",
        "Python * Documentation*",
        "Python * Executables*",
        "Python * pip Bootstrap*",
        "Python * Standard Library*",
        "Python * Tcl/Tk Support*",
        "Python * Test Suite*",
        "Python * Utility Scripts*",        # ajout validé
        # « Python Launcher » est CONSERVÉ (décision de portée).
    ],
    "Autres": [
        "Microsoft PowerAutomateDesktop",
        "Microsoft Windows Application Compatibility Fix Database",  # ajout validé
        "Microsoft WindowsScan",
        "Microsoft MicrosoftJournal",
        "SnippingTool-ee6eb196-db28-4d99-816d-fa9a63b4a377",
        "TechSmith Screen Capture Codec",
    ],
}

_MOTIFS_COMPILES = [
    (categorie, motif, _compiler_motif(motif))
    for categorie, motifs in _MOTIFS.items()
    for motif in motifs
]


def _trouver_motif(logiciel):
    """Retourne (catégorie, motif) du 1er motif correspondant, sinon None."""
    texte = _norm(logiciel)
    if not texte:
        return None
    for categorie, motif, regex in _MOTIFS_COMPILES:
        if regex.match(texte):
            return (categorie, motif)
    return None


def _analyser_lignes(lignes, cols):
    """Identifie les lignes à exclure parmi `lignes` ((index, valeurs)).

    `cols` provient de colonnes.index_colonnes() : la colonne « logiciel »
    est résolue par son LIBELLÉ d'en-tête, jamais par une position fixe.

    Retourne (indices, correspondances) où :
    - indices : liste des index de lignes à déplacer;
    - correspondances : {(catégorie, motif): {libellé: nb de lignes}}
      pour le rapport détaillé.
    """
    indices = []
    correspondances = {}
    idx_log = cols["logiciel"]
    for idx, row in lignes:
        logiciel = valeur(row, idx_log)
        trouve = _trouver_motif(logiciel)
        if trouve is None:
            continue
        indices.append(idx)
        libelle = str(logiciel).strip()
        par_libelle = correspondances.setdefault(trouve, {})
        par_libelle[libelle] = par_libelle.get(libelle, 0) + 1
    return indices, correspondances


def _secteurs(sheetnames):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement."""
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in sheetnames:
        if name.casefold() not in ignored:
            secteurs.append(name)
        if name.casefold() == LAST_SECTOR.casefold():
            break
    return secteurs


def _fmt_duree(secondes):
    """Formate une durée en secondes sous la forme « 1 h 02 min 03 s »."""
    secondes = int(secondes)
    h, reste = divmod(secondes, 3600)
    m, s = divmod(reste, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"


def main():
    # Évite les UnicodeEncodeError quand la sortie est redirigée vers un
    # fichier (encodage console cp1252) : les caractères non représentables
    # sont remplacés plutôt que de faire planter le script.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(
        description="Applique la Règle 6 (exclusions spécifiques de "
                    "composantes Autodesk, Adobe, Citrix, Python et autres).")
    parser.add_argument("--dry-run", action="store_true",
                        help="Simulation : affiche les résultats sans modifier le classeur.")
    args = parser.parse_args()

    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
        shutil.copy2(WORKBOOK, backup)
        print(f"Copie de sauvegarde : {backup.name}", flush=True)

    # Étape 1 — analyse en lecture seule (rapide, peu de mémoire) :
    # liste détaillée des applications correspondantes et nombre de lignes
    # à déplacer par secteur. Aucun déplacement à cette étape.
    print("\nÉtape 1/2 — Analyse des onglets (lecture seule)...", flush=True)
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    try:
        secteurs = _secteurs(wb.sheetnames)
        print(f"Onglets de secteurs à traiter (jusqu'à « {LAST_SECTOR} ») : "
              f"{len(secteurs)}", flush=True)
        header = None
        compteurs = {}           # {onglet: nb de lignes à déplacer}
        corresp_par_onglet = {}  # {onglet: {(catégorie, motif): {libellé: nb}}}
        entetes = {}             # {onglet: en-tête d'origine}
        index_par_onglet = {}    # {onglet: index des colonnes requises}
        for sheet_name in secteurs:
            ws = wb[sheet_name]
            lignes_it = ws.iter_rows(values_only=True)
            entete = next(lignes_it, None)
            if entete is not None:
                entetes[sheet_name] = list(entete)
                # La colonne « logiciel » est localisée par son LIBELLÉ,
                # jamais par une position fixe (voir colonnes.py).
                index_par_onglet[sheet_name] = index_colonnes(
                    entete, ("logiciel",), onglet=sheet_name)
                if header is None:
                    header = list(entete)
            lignes = list(enumerate(lignes_it, start=2))
            indices, corresp = _analyser_lignes(lignes,
                                                index_par_onglet[sheet_name])
            compteurs[sheet_name] = len(indices)
            corresp_par_onglet[sheet_name] = corresp
    finally:
        wb.close()
        del wb
        gc.collect()

    # Regroupe les correspondances de tous les onglets pour le rapport.
    globales = {}  # {(catégorie, motif): {libellé: nb de lignes total}}
    for corresp in corresp_par_onglet.values():
        for cle, par_libelle in corresp.items():
            dest = globales.setdefault(cle, {})
            for libelle, nb in par_libelle.items():
                dest[libelle] = dest.get(libelle, 0) + nb

    print("\nListe détaillée des applications exclues "
          "(par catégorie et motif) :")
    for categorie, motifs in _MOTIFS.items():
        print(f"\n  Catégorie « {categorie} » :")
        for motif in motifs:
            par_libelle = globales.get((categorie, motif))
            if not par_libelle:
                print(f"    Motif « {motif} » : aucune correspondance")
                continue
            total = sum(par_libelle.values())
            print(f"    Motif « {motif} » : {total} ligne(s), "
                  f"{len(par_libelle)} application(s) distincte(s)")
            for libelle in sorted(par_libelle, key=str.casefold):
                print(f"      « {libelle} » — {par_libelle[libelle]} ligne(s)")

    total_planned = sum(compteurs.values())
    print("\nLignes à déplacer par onglet :")
    for sheet_name in secteurs:
        if compteurs[sheet_name]:
            print(f"  {sheet_name} : {compteurs[sheet_name]}")
    print(f"\nLignes à déplacer (tous secteurs) : {total_planned}", flush=True)
    if args.dry_run:
        print("Mode --dry-run : aucune modification.")
        return

    # Étape 2 — un secteur à la fois : ouverture du classeur, déplacement des
    # lignes correspondantes vers « Excluded Regle06 », sauvegarde, fermeture
    # et libération des ressources avant de passer au secteur suivant
    # (aucune pause).
    print(f"\nÉtape 2/2 — Traitement des secteurs (un à la fois, "
          f"sauvegarde après chacun) :", flush=True)
    debut = time.perf_counter()
    total_deplace = 0
    nb_secteurs = len(secteurs)
    for i, sheet_name in enumerate(secteurs, start=1):
        ecoule = time.perf_counter() - debut
        if i > 1:
            restant = ecoule / (i - 1) * (nb_secteurs - i + 1)
            progression = (f"temps écoulé {_fmt_duree(ecoule)}, "
                           f"restant estimé {_fmt_duree(restant)}")
        else:
            progression = "estimation du temps restant après le 1er onglet"
        print(f"\n[{i}/{nb_secteurs}] Onglet « {sheet_name} » ({progression})",
              flush=True)

        if compteurs[sheet_name] == 0:
            print("  -> 0 ligne à déplacer, onglet inchangé.", flush=True)
            continue

        t0 = time.perf_counter()
        wb = openpyxl.load_workbook(WORKBOOK)
        try:
            ws = wb[sheet_name]
            lignes = list(enumerate(ws.iter_rows(min_row=2, values_only=True),
                                    start=2))
            indices, _corresp = _analyser_lignes(lignes,
                                                 index_par_onglet[sheet_name])
            valeurs_par_idx = {idx: row for idx, row in lignes}
            matches = [(idx, valeurs_par_idx[idx]) for idx in indices]

            excluded_ws = (wb[EXCLUDED_SHEET]
                           if EXCLUDED_SHEET in wb.sheetnames else None)
            if excluded_ws is None:
                excluded_ws = wb.create_sheet(EXCLUDED_SHEET)
                excluded_ws.append((header or []) + ["Raison"])
            for _row_idx, values in matches:
                # Réordonne la ligne vers l'en-tête de référence : les onglets
                # de secteurs peuvent présenter un ordre de colonnes différent
                # sans que l'onglet d'exclusion ne scramblie les colonnes.
                excluded_ws.append(
                    aligner(values, entetes[sheet_name], header,
                            onglet=sheet_name) + [REASON])
            # Suppression par blocs de lignes consécutives (de bas en haut) :
            # beaucoup plus rapide qu'une suppression ligne par ligne.
            indices = sorted(indices)
            i = len(indices) - 1
            while i >= 0:
                j = i
                while j > 0 and indices[j - 1] == indices[j] - 1:
                    j -= 1
                ws.delete_rows(indices[j], indices[i] - indices[j] + 1)
                i = j - 1

            # Ajuste les plages des tableaux structurés et des filtres à la
            # nouvelle fin de données. Sans cela, les plages pointent au-delà
            # des lignes restantes et Excel signale un contenu corrompu à
            # l'ouverture du fichier.
            fin = max(ws.max_row, 1)
            for tbl in ws.tables.values():
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", tbl.ref)
                if m and int(m.group(3)) != fin:
                    nouveau = f"{m.group(1)}:{m.group(2)}{fin}"
                    tbl.ref = nouveau
                    if tbl.autoFilter is not None:
                        tbl.autoFilter.ref = nouveau
            if ws.auto_filter.ref:
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", ws.auto_filter.ref)
                if m and int(m.group(3)) != fin:
                    ws.auto_filter.ref = f"{m.group(1)}:{m.group(2)}{fin}"

            wb.save(WORKBOOK)
        finally:
            wb.close()
            del wb
            gc.collect()  # libère les ressources avant le secteur suivant

        total_deplace += len(matches)
        print(f"  -> {len(matches)} ligne(s) déplacée(s) et sauvegardée(s) "
              f"en {_fmt_duree(time.perf_counter() - t0)}.", flush=True)

    print(f"\nTerminé en {_fmt_duree(time.perf_counter() - debut)}.")
    print(f"Total déplacé : {total_deplace} ligne(s) -> « {EXCLUDED_SHEET} ».")
    print("Aucune perte de données : chaque ligne retirée y est copiée telle quelle.")


if __name__ == "__main__":
    main()
