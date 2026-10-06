"""
appliquer_regle02.py — Applique la Règle 2 (exclusion des applications
Microsoft standards faisant partie du poste standard ou de l'environnement
de base) sur les onglets de secteurs de
« User_Workstation_App_Associations_Consolidated.xlsx »
(jusqu'à l'onglet « Non défini » inclusivement).

Comportement (conforme à Regle_02_microsoft_standards.md) :
1. Génère d'abord la liste distincte des applications correspondant à la
   règle (aucun déplacement). Cette analyse est faite en lecture seule
   (read_only) pour être rapide et économe en mémoire.
2. Traite ensuite UN secteur à la fois : le classeur est rouvert pour cet
   onglet seulement, les lignes correspondantes sont déplacées vers l'onglet
   « Excluded Regle02 » (créé une seule fois) sans transformation des données
   d'origine (une colonne finale « Raison » est ajoutée dans l'onglet
   d'exclusion), puis le classeur est sauvegardé, fermé et les ressources
   sont libérées avant de passer au secteur suivant.
3. La progression est affichée à l'écran (onglet courant, temps écoulé et
   estimation du temps restant). Aucune intervention n'est requise entre
   les secteurs.

Sont volontairement CONSERVÉES (garde-fou) les applications Microsoft qui ne
font pas partie de la suite Office de base : Visio, Power BI, Visual Studio
(Code), Microsoft Project, ainsi que Microsoft 365 Copilot.

Usage : py appliquer_regle02.py --dry-run   # simulation (aucune écriture)
        py appliquer_regle02.py             # exécution réelle
"""

import argparse
import gc
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from colonnes import aligner, index_colonnes, valeur

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
EXCLUDED_SHEET = "Excluded Regle02"
REASON = "Règle 2 — Application Microsoft standard (poste standard / environnement de base)"
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
IGNORED = {"logiciels", "sommaire", "feuil1", "excluded regle01", EXCLUDED_SHEET.casefold()}


def _norm(value):
    """Normalise une valeur : casefold, apostrophes unifiées, espaces réduits."""
    text = "" if value is None else str(value)
    for ch in ("’", "‘", "`"):
        text = text.replace(ch, "'")
    return " ".join(text.split()).casefold()


# --- Garde-fou : applications Microsoft hors suite de base (à CONSERVER) ----
# Vérifié avant toute catégorie : si l'un de ces motifs est présent, la ligne
# n'est jamais exclue par la Règle 2 (conforme au MD : Visio, Power BI,
# Visual Studio Code; Microsoft 365 Copilot conservé explicitement).
_KEEP_SUBSTRINGS = (
    "microsoft 365 copilot",   # conservé explicitement par le MD
    "visio",
    "power bi",
    "powerbi",
    "visual studio",
    "microsoft project",
)

# --- A) Suite Office 365 / Microsoft 365 et composantes Microsoft Office ----
_OFFICE_SUBSTRINGS = (
    "microsoft 365",           # Microsoft 365 - fr-fr, Apps for business/enterprise,
                               # companion apps, Aplicaciones... (sauf 365 Copilot, voir garde-fou)
    "office 365",
    "microsoft office",        # Office 2016 (Pro Plus, MUI, Proofing, OSM,
                               # Access database engine, outils linguistiques FR/ES/DE/NL)
    "office 16 click-to-run",  # Office 16 Click-to-Run Licensing Component
    "microsoft word",          # Microsoft Word MUI (French) 2016
    "microsoft excel",         # Microsoft Excel MUI (French) 2016
    "microsoft powerpoint",    # Microsoft PowerPoint MUI (French) 2016
    "microsoft outlook",       # Microsoft Outlook MUI (French) 2016
    "outlook for windows",
    "skype for business",      # composante Office (ex. MUI French 2016)
)

# --- B) Microsoft Edge -------------------------------------------------------
_EDGE_SUBSTRINGS = (
    "microsoft edge",          # Microsoft Edge, Assistant de jeu Microsoft Edge
    "microsoftedge",           # Microsoft MicrosoftEdge
)

# --- C) OneDrive -------------------------------------------------------------
_ONEDRIVE_SUBSTRINGS = (
    "microsoft onedrive",
)
_ONEDRIVE_EXACT = {
    "onedrive",                # « HP OneDrive Plugin » (tiers) est conservé
}

# --- D) OneNote ----------------------------------------------------------------
_ONENOTE_SUBSTRINGS = (
    "onenote",                 # Microsoft OneNote - fr-fr/en-us/es-es, MUI 2016,
                               # OneNote for Windows 10
)

# --- E) Microsoft Teams --------------------------------------------------------
_TEAMS_SUBSTRINGS = (
    "microsoft teams",         # Microsoft Teams, Microsoft Teams VDI Citrix plugin
    "teams machine-wide installer",
    # « kuando Busylight for MS Teams » (tiers) est conservé
)

# --- F) Microsoft Whiteboard ---------------------------------------------------
_WHITEBOARD_SUBSTRINGS = (
    "whiteboard",
)

# --- G) Copilot (sauf Microsoft 365 Copilot, voir garde-fou) -------------------
_COPILOT_SUBSTRINGS = (
    "copilot",                 # Copilot, Windows CoPilot MSIX Pack
)

# --- H) Bing / Exchange / Loop / Stream ----------------------------------------
_BING_SUBSTRINGS = (
    "bing",                    # Bing Wallpaper, Microsoft BingNews/BingTranslator/BingWeather
)
_MS_LOOP_STREAM_SUBSTRINGS = (
    "microsoft loop",
    "microsoft stream",        # « NVIDIA SHIELD Streaming » et « PaperStream » non touchés
)
_EXCHANGE_RE = re.compile(r"\bexchange\b")

# --- I) Autres applications Microsoft standards (grand public) -----------------
_EXTRAS_SUBSTRINGS = (
    "clipchamp",               # Clipchamp, Clipchamp.Clipchamp
    "fresh paint",
)
_SWAY_RE = re.compile(r"\bsway\b")     # Sway


# --- J) Composantes et applis Microsoft in-box / agents (ajouts) ---------------
_MS_INBOX_SUBSTRINGS = (
    "microsoft epm agent",     # Microsoft EPM Agent
    "devhome",                 # Microsoft Windows DevHome
    "dev home",                # Dev Home (Preview)
    "soundrecorder",           # Microsoft WindowsSoundRecorder
    "xbox",                    # Microsoft XboxApp, XboxIdentityProvider, XboxSpeechToTextOverlay
    "yourphone",               # Microsoft YourPhone (Phone Link)
    "zune",                    # Microsoft ZuneMusic, Microsoft ZuneVideo
)



def match_rule2(logiciel, manufacturier):
    """Retourne la catégorie ('A' à 'J') si la ligne correspond à la
    Règle 2, sinon None."""
    name = _norm(logiciel)
    if not name:
        return None

    # Garde-fou : applications Microsoft hors suite de base, toujours conservées.
    if any(s in name for s in _KEEP_SUBSTRINGS):
        return None

    if any(s in name for s in _OFFICE_SUBSTRINGS):
        return "A"
    if any(s in name for s in _EDGE_SUBSTRINGS):
        return "B"
    if any(s in name for s in _ONEDRIVE_SUBSTRINGS) or name in _ONEDRIVE_EXACT:
        return "C"
    if any(s in name for s in _ONENOTE_SUBSTRINGS):
        return "D"
    if any(s in name for s in _TEAMS_SUBSTRINGS):
        return "E"
    if any(s in name for s in _WHITEBOARD_SUBSTRINGS):
        return "F"
    if any(s in name for s in _COPILOT_SUBSTRINGS):
        return "G"
    if (any(s in name for s in _BING_SUBSTRINGS)
            or any(s in name for s in _MS_LOOP_STREAM_SUBSTRINGS)
            or _EXCHANGE_RE.search(name)):
        return "H"
    if (any(s in name for s in _EXTRAS_SUBSTRINGS)
            or _SWAY_RE.search(name)):
        return "I"
    if any(s in name for s in _MS_INBOX_SUBSTRINGS):
        return "J"
    return None


def _secteurs(sheetnames):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement."""
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in sheetnames:
        if name.casefold() not in ignored:
            secteurs.append(name)
        if name.casefold() == LAST_SECTOR.casefold():
            break
    return secteurs


def _fmt_duree(secondes):
    """Formate une durée en secondes sous la forme « 1 h 02 min 03 s »."""
    secondes = int(secondes)
    h, reste = divmod(secondes, 3600)
    m, s = divmod(reste, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"


def main():
    # Évite les UnicodeEncodeError quand la sortie est redirigée vers un
    # fichier (encodage console cp1252) : les caractères non représentables
    # sont remplacés plutôt que de faire planter le script.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(
        description="Applique la Règle 2 (applications Microsoft standards).")
    parser.add_argument("--dry-run", action="store_true",
                        help="Simulation : affiche les résultats sans modifier le classeur.")
    args = parser.parse_args()

    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
        shutil.copy2(WORKBOOK, backup)
        print(f"Copie de sauvegarde : {backup.name}", flush=True)

    # Étape 1 — analyse en lecture seule (rapide, peu de mémoire) :
    # liste distincte des applications correspondant à la règle et nombre
    # de lignes à déplacer par secteur. Aucun déplacement à cette étape.
    print("\nÉtape 1/2 — Analyse des onglets (lecture seule)...", flush=True)
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    try:
        secteurs = _secteurs(wb.sheetnames)
        print(f"Onglets de secteurs à traiter (jusqu'à « {LAST_SECTOR} ») : "
              f"{len(secteurs)}", flush=True)
        header = None
        distinct = set()
        compteurs = {}   # {onglet: nb de lignes correspondantes}
        entetes = {}     # {onglet: en-tête d'origine}
        index_par_onglet = {}  # {onglet: index des colonnes requises}
        for sheet_name in secteurs:
            ws = wb[sheet_name]
            lignes = ws.iter_rows(values_only=True)
            entete = next(lignes, None)
            if entete is not None:
                entetes[sheet_name] = list(entete)
                # Les colonnes sont localisées par leur LIBELLÉ, jamais par
                # leur position : l'ordre peut varier d'une extraction à
                # l'autre. Lève une RuntimeError explicite si une colonne
                # requise est absente.
                index_par_onglet[sheet_name] = index_colonnes(
                    entete, ("logiciel", "manufacturier"), onglet=sheet_name)
                if header is None:
                    header = list(entete)
            cols = index_par_onglet[sheet_name]
            count = 0
            for row in lignes:
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                categorie = match_rule2(logiciel, manufacturier)
                if categorie:
                    count += 1
                    distinct.add((str(logiciel).strip(), categorie))
            compteurs[sheet_name] = count
    finally:
        wb.close()
        del wb
        gc.collect()

    print("\nListe distincte des applications correspondant à la Règle 2 :")
    for logiciel, categorie in sorted(distinct, key=lambda t: (t[1], t[0].casefold())):
        print(f"  [{categorie}] {logiciel}")
    total_planned = sum(compteurs.values())
    print(f"\nApplications distinctes : {len(distinct)}")
    print(f"Lignes à déplacer (tous secteurs) : {total_planned}", flush=True)
    if args.dry_run:
        print("Mode --dry-run : aucune modification.")
        return



    # Étape 2 — un secteur à la fois : ouverture du classeur, déplacement des
    # lignes vers « Excluded Regle02 », sauvegarde, fermeture et libération
    # des ressources avant de passer au secteur suivant (aucune pause).
    print(f"\nÉtape 2/2 — Traitement des secteurs (un à la fois, "
          f"sauvegarde après chacun) :", flush=True)
    debut = time.perf_counter()
    total_deplace = 0
    nb_secteurs = len(secteurs)
    for i, sheet_name in enumerate(secteurs, start=1):
        ecoule = time.perf_counter() - debut
        if i > 1:
            restant = ecoule / (i - 1) * (nb_secteurs - i + 1)
            progression = (f"temps écoulé {_fmt_duree(ecoule)}, "
                           f"restant estimé {_fmt_duree(restant)}")
        else:
            progression = "estimation du temps restant après le 1er onglet"
        print(f"\n[{i}/{nb_secteurs}] Onglet « {sheet_name} » ({progression})",
              flush=True)

        if compteurs[sheet_name] == 0:
            print("  -> 0 ligne à déplacer, onglet inchangé.", flush=True)
            continue

        t0 = time.perf_counter()
        wb = openpyxl.load_workbook(WORKBOOK)
        try:
            ws = wb[sheet_name]
            cols = index_par_onglet[sheet_name]
            matches = []
            for row_idx, row in enumerate(
                    ws.iter_rows(min_row=2, values_only=True), start=2):
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                if match_rule2(logiciel, manufacturier):
                    matches.append((row_idx, row))

            excluded_ws = (wb[EXCLUDED_SHEET]
                           if EXCLUDED_SHEET in wb.sheetnames else None)
            if excluded_ws is None:
                excluded_ws = wb.create_sheet(EXCLUDED_SHEET)
                excluded_ws.append((header or []) + ["Raison"])
            for _row_idx, values in matches:
                # Réordonne la ligne vers l'en-tête de référence : les onglets
                # de secteurs peuvent présenter un ordre de colonnes différent
                # sans que l'onglet d'exclusion ne scramblie les colonnes.
                excluded_ws.append(
                    aligner(values, entetes[sheet_name], header,
                            onglet=sheet_name) + [REASON])
            # Suppression par blocs de lignes consécutives (de bas en haut) :
            # beaucoup plus rapide qu'une suppression ligne par ligne.
            indices = sorted(m[0] for m in matches)
            i = len(indices) - 1
            while i >= 0:
                j = i
                while j > 0 and indices[j - 1] == indices[j] - 1:
                    j -= 1
                ws.delete_rows(indices[j], indices[i] - indices[j] + 1)
                i = j - 1

            # Ajuste les plages des tableaux structurés et des filtres à la
            # nouvelle fin de données. Sans cela, les plages pointent au-delà
            # des lignes restantes et Excel signale un contenu corrompu à
            # l'ouverture du fichier.
            fin = max(ws.max_row, 1)
            for tbl in ws.tables.values():
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", tbl.ref)
                if m and int(m.group(3)) != fin:
                    nouveau = f"{m.group(1)}:{m.group(2)}{fin}"
                    tbl.ref = nouveau
                    if tbl.autoFilter is not None:
                        tbl.autoFilter.ref = nouveau
            if ws.auto_filter.ref:
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", ws.auto_filter.ref)
                if m and int(m.group(3)) != fin:
                    ws.auto_filter.ref = f"{m.group(1)}:{m.group(2)}{fin}"

            wb.save(WORKBOOK)
        finally:
            wb.close()
            del wb
            gc.collect()  # libère les ressources avant le secteur suivant

        total_deplace += len(matches)
        print(f"  -> {len(matches)} ligne(s) déplacée(s) et sauvegardée(s) "
              f"en {_fmt_duree(time.perf_counter() - t0)}.", flush=True)

    print(f"\nTerminé en {_fmt_duree(time.perf_counter() - debut)}.")
    print(f"Total déplacé : {total_deplace} ligne(s) -> « {EXCLUDED_SHEET} ».")
    print("Aucune perte de données : chaque ligne retirée y est copiée telle quelle.")


if __name__ == "__main__":
    main()
