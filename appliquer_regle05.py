"""
appliquer_regle05.py — Applique la Règle 5 (dédoublonnage des applications
entre plusieurs sources) sur les onglets de secteurs de
« User_Workstation_App_Associations_Consolidated.xlsx »
(jusqu'à l'onglet « Non défini » inclusivement).

Règle (conforme à Regle_05_dedoublonnage_sources.md) :
lorsque la même application est présente pour le même « Nom d'hôte »,
conserver une seule occurrence; les occurrences excédentaires sont déplacées
vers l'onglet « Excluded Regle05 » sans transformation (aucune perte de
données — une colonne finale « Raison » est ajoutée dans l'onglet
d'exclusion).

Détection d'un doublon (à l'intérieur de chaque onglet de secteur
seulement — les doublons répartis entre deux onglets sont conservés,
décision de portée) :
- clé = (Logiciel normalisé, Nom d'hôte normalisé);
- normalisation poussée du logiciel : casse ignorée, apostrophes unifiées,
  tirets/tirets cadratin unifiés, puis retrait de tout caractère non
  alphanumérique (espaces, ponctuation). Ainsi « Microsoft Power BI
  Desktop (x64) » et « Microsoft PowerBI Desktop (x64) » sont des doublons;
- normalisation de l'hôte : casse ignorée, espaces superflus retirés
  (« cnh-0186 » = « CNH-0186 »).

Occurrence conservée :
- si un libellé canonique est défini pour la clé (table
  _LIBELLES_CANONIQUES), on conserve la ligne portant ce libellé
  (ex. « Microsoft Power BI Desktop (x64) » avec espace);
- sinon, la première occurrence rencontrée dans l'onglet.

Comportement (calqué sur appliquer_regle04.py) :
1. Génère d'abord la liste détaillée des groupes de doublons (aucun
   déplacement), en lecture seule (read_only) pour être rapide et économe
   en mémoire. Ce rapport permet de valider les variantes interprétées
   (libellés non identiques) avant l'exécution réelle.
2. Traite ensuite UN secteur à la fois : le classeur est rouvert pour cet
   onglet seulement, les lignes excédentaires sont déplacées vers l'onglet
   « Excluded Regle05 » (créé une seule fois), puis le classeur est
   sauvegardé, fermé et les ressources sont libérées avant de passer au
   secteur suivant.
3. La progression est affichée à l'écran (onglet courant, temps écoulé et
   estimation du temps restant). Aucune intervention n'est requise entre
   les secteurs.

Usage : py appliquer_regle05.py --dry-run   # simulation (aucune écriture)
        py appliquer_regle05.py             # exécution réelle
"""

import argparse
import gc
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from colonnes import aligner, index_colonnes, valeur

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
EXCLUDED_SHEET = "Excluded Regle05"
REASON = "Règle 5 — Doublon : même application, même nom d'hôte"
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
IGNORED = {"logiciels", "sommaire", "feuil1", "excluded regle01", "excluded regle02",
           "excluded regle03", "excluded regle04", EXCLUDED_SHEET.casefold()}


def _norm(value):
    """Normalise une valeur : casefold, apostrophes unifiées, espaces réduits."""
    text = "" if value is None else str(value)
    for ch in ("’", "‘", "`"):
        text = text.replace(ch, "'")
    return " ".join(text.split()).casefold()


def _cle_logiciel(value):
    """Clé de dédoublonnage du logiciel : normalisation poussée.

    Casse ignorée, apostrophes et tirets (simple, cadratin, demi-cadratin,
    moins) unifiés, puis retrait de tout caractère non alphanumérique
    (espaces, ponctuation). Les accents sont conservés (« é » reste « é »).
    """
    text = _norm(value)
    for ch in ("–", "—", "−"):
        text = text.replace(ch, "-")
    return "".join(ch for ch in text if ch.isalnum())


# Libellés canoniques : pour les variantes d'un même logiciel, la ligne
# portant ce libellé est conservée de préférence à la première occurrence.
# Clé = _cle_logiciel(libellé), valeur = libellé exact à conserver.
_LIBELLES_CANONIQUES = {
    _cle_logiciel("Microsoft Power BI Desktop (x64)"):
        "Microsoft Power BI Desktop (x64)",
}



def _analyser_lignes(lignes, cols):
    """Identifie les doublons parmi `lignes` (liste de (index, valeurs)).

    `cols` provient de colonnes.index_colonnes() : les colonnes « logiciel »
    et « hôte » sont résolues par leur LIBELLÉ d'en-tête, jamais par une
    position fixe. Une position fixe ferait dédupliquer sur la mauvaise
    colonne et supprimerait des applications légitimes si l'ordre des
    colonnes change.

    Retourne (indices_a_deplacer, groupes) où :
    - indices_a_deplacer : liste des index de lignes excédentaires;
    - groupes : liste de dicts décrivant chaque groupe de doublons
      (hôte, libellés, libellé conservé, nombre déplacé) pour le rapport.
    """
    idx_log = cols["logiciel"]
    idx_hote = cols["hote"]
    par_cle = {}
    for idx, row in lignes:
        logiciel = valeur(row, idx_log)
        hote = valeur(row, idx_hote)
        cle_log = _cle_logiciel(logiciel)
        cle_hote = _norm(hote)
        if not cle_log or not cle_hote:
            continue  # ligne sans logiciel ou sans hôte : jamais un doublon
        par_cle.setdefault((cle_log, cle_hote), []).append((idx, row))

    indices = []
    groupes = []
    for (cle_log, cle_hote), occurrences in par_cle.items():
        if len(occurrences) < 2:
            continue
        canonique = _LIBELLES_CANONIQUES.get(cle_log)
        garde = None
        if canonique is not None:
            for idx, row in occurrences:
                if str(valeur(row, idx_log)).strip() == canonique:
                    garde = idx
                    break
        if garde is None:
            garde = occurrences[0][0]
        deplaces = [idx for idx, _row in occurrences if idx != garde]
        indices.extend(deplaces)
        libelle_garde = next(str(valeur(row, idx_log)).strip()
                             for idx, row in occurrences if idx == garde)
        groupes.append({
            "hote": str(valeur(occurrences[0][1], idx_hote)).strip(),
            "libelles": sorted({str(valeur(row, idx_log)).strip()
                                for _idx, row in occurrences}),
            "garde": libelle_garde,
            "nb_occurrences": len(occurrences),
            "nb_deplaces": len(deplaces),
        })
    return indices, groupes


def _secteurs(sheetnames):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement."""
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in sheetnames:
        if name.casefold() not in ignored:
            secteurs.append(name)
        if name.casefold() == LAST_SECTOR.casefold():
            break
    return secteurs


def _fmt_duree(secondes):
    """Formate une durée en secondes sous la forme « 1 h 02 min 03 s »."""
    secondes = int(secondes)
    h, reste = divmod(secondes, 3600)
    m, s = divmod(reste, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"


def main():
    # Évite les UnicodeEncodeError quand la sortie est redirigée vers un
    # fichier (encodage console cp1252) : les caractères non représentables
    # sont remplacés plutôt que de faire planter le script.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(
        description="Applique la Règle 5 (dédoublonnage des applications "
                    "pour un même nom d'hôte).")
    parser.add_argument("--dry-run", action="store_true",
                        help="Simulation : affiche les résultats sans modifier le classeur.")
    args = parser.parse_args()

    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
        shutil.copy2(WORKBOOK, backup)
        print(f"Copie de sauvegarde : {backup.name}", flush=True)

    # Étape 1 — analyse en lecture seule (rapide, peu de mémoire) :
    # liste détaillée des groupes de doublons et nombre de lignes à déplacer
    # par secteur. Aucun déplacement à cette étape.
    print("\nÉtape 1/2 — Analyse des onglets (lecture seule)...", flush=True)
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    try:
        secteurs = _secteurs(wb.sheetnames)
        print(f"Onglets de secteurs à traiter (jusqu'à « {LAST_SECTOR} ») : "
              f"{len(secteurs)}", flush=True)
        header = None
        compteurs = {}           # {onglet: nb de lignes à déplacer}
        groupes_par_onglet = {}  # {onglet: [groupes de doublons]}
        entetes = {}             # {onglet: en-tête d'origine}
        index_par_onglet = {}    # {onglet: index des colonnes requises}
        for sheet_name in secteurs:
            ws = wb[sheet_name]
            lignes_it = ws.iter_rows(values_only=True)
            entete = next(lignes_it, None)
            if entete is not None:
                entetes[sheet_name] = list(entete)
                # « logiciel » et « hôte » sont localisés par leur LIBELLÉ,
                # jamais par une position fixe (voir colonnes.py) : sinon la
                # déduplication porterait sur la mauvaise colonne.
                index_par_onglet[sheet_name] = index_colonnes(
                    entete, ("logiciel", "hote"), onglet=sheet_name)
                if header is None:
                    header = list(entete)
            lignes = list(enumerate(lignes_it, start=2))
            indices, groupes = _analyser_lignes(lignes,
                                                index_par_onglet[sheet_name])
            compteurs[sheet_name] = len(indices)
            groupes_par_onglet[sheet_name] = groupes
    finally:
        wb.close()
        del wb
        gc.collect()

    print("\nListe détaillée des groupes de doublons "
          "(même application, même nom d'hôte, par onglet) :")
    for sheet_name in secteurs:
        groupes = groupes_par_onglet[sheet_name]
        if not groupes:
            continue
        print(f"\n  Onglet « {sheet_name} » :")
        for g in sorted(groupes, key=lambda g: (g["hote"].casefold(),
                                                g["garde"].casefold())):
            print(f"    Hôte {g['hote']} — {g['nb_occurrences']} occurrence(s), "
                  f"1 conservée (« {g['garde']} »), "
                  f"{g['nb_deplaces']} déplacée(s)")
            if len(g["libelles"]) > 1:
                variantes = " », « ".join(g["libelles"])
                print(f"      variantes interprétées comme doublons : "
                      f"« {variantes} »")
    total_planned = sum(compteurs.values())
    nb_groupes = sum(len(g) for g in groupes_par_onglet.values())
    print(f"\nGroupes de doublons : {nb_groupes}")
    print(f"Lignes à déplacer (tous secteurs) : {total_planned}", flush=True)
    if args.dry_run:
        print("Mode --dry-run : aucune modification.")
        return

    # Étape 2 — un secteur à la fois : ouverture du classeur, déplacement des
    # lignes excédentaires vers « Excluded Regle05 », sauvegarde, fermeture
    # et libération des ressources avant de passer au secteur suivant
    # (aucune pause).
    print(f"\nÉtape 2/2 — Traitement des secteurs (un à la fois, "
          f"sauvegarde après chacun) :", flush=True)
    debut = time.perf_counter()
    total_deplace = 0
    nb_secteurs = len(secteurs)
    for i, sheet_name in enumerate(secteurs, start=1):
        ecoule = time.perf_counter() - debut
        if i > 1:
            restant = ecoule / (i - 1) * (nb_secteurs - i + 1)
            progression = (f"temps écoulé {_fmt_duree(ecoule)}, "
                           f"restant estimé {_fmt_duree(restant)}")
        else:
            progression = "estimation du temps restant après le 1er onglet"
        print(f"\n[{i}/{nb_secteurs}] Onglet « {sheet_name} » ({progression})",
              flush=True)

        if compteurs[sheet_name] == 0:
            print("  -> 0 ligne à déplacer, onglet inchangé.", flush=True)
            continue

        t0 = time.perf_counter()
        wb = openpyxl.load_workbook(WORKBOOK)
        try:
            ws = wb[sheet_name]
            lignes = list(enumerate(ws.iter_rows(min_row=2, values_only=True),
                                    start=2))
            indices, _groupes = _analyser_lignes(lignes,
                                                 index_par_onglet[sheet_name])
            valeurs_par_idx = {idx: row for idx, row in lignes}
            matches = [(idx, valeurs_par_idx[idx]) for idx in indices]

            excluded_ws = (wb[EXCLUDED_SHEET]
                           if EXCLUDED_SHEET in wb.sheetnames else None)
            if excluded_ws is None:
                excluded_ws = wb.create_sheet(EXCLUDED_SHEET)
                excluded_ws.append((header or []) + ["Raison"])
            for _row_idx, values in matches:
                # Réordonne la ligne vers l'en-tête de référence : les onglets
                # de secteurs peuvent présenter un ordre de colonnes différent
                # sans que l'onglet d'exclusion ne scramblie les colonnes.
                excluded_ws.append(
                    aligner(values, entetes[sheet_name], header,
                            onglet=sheet_name) + [REASON])
            # Suppression par blocs de lignes consécutives (de bas en haut) :
            # beaucoup plus rapide qu'une suppression ligne par ligne.
            indices = sorted(indices)
            i = len(indices) - 1
            while i >= 0:
                j = i
                while j > 0 and indices[j - 1] == indices[j] - 1:
                    j -= 1
                ws.delete_rows(indices[j], indices[i] - indices[j] + 1)
                i = j - 1

            # Ajuste les plages des tableaux structurés et des filtres à la
            # nouvelle fin de données. Sans cela, les plages pointent au-delà
            # des lignes restantes et Excel signale un contenu corrompu à
            # l'ouverture du fichier.
            fin = max(ws.max_row, 1)
            for tbl in ws.tables.values():
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", tbl.ref)
                if m and int(m.group(3)) != fin:
                    nouveau = f"{m.group(1)}:{m.group(2)}{fin}"
                    tbl.ref = nouveau
                    if tbl.autoFilter is not None:
                        tbl.autoFilter.ref = nouveau
            if ws.auto_filter.ref:
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", ws.auto_filter.ref)
                if m and int(m.group(3)) != fin:
                    ws.auto_filter.ref = f"{m.group(1)}:{m.group(2)}{fin}"

            wb.save(WORKBOOK)
        finally:
            wb.close()
            del wb
            gc.collect()  # libère les ressources avant le secteur suivant

        total_deplace += len(matches)
        print(f"  -> {len(matches)} ligne(s) déplacée(s) et sauvegardée(s) "
              f"en {_fmt_duree(time.perf_counter() - t0)}.", flush=True)

    print(f"\nTerminé en {_fmt_duree(time.perf_counter() - debut)}.")
    print(f"Total déplacé : {total_deplace} ligne(s) -> « {EXCLUDED_SHEET} ».")
    print("Aucune perte de données : chaque ligne retirée y est copiée telle quelle.")


if __name__ == "__main__":
    main()


