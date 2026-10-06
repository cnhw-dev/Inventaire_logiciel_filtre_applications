"""
appliquer_regle03.py — Applique la Règle 3 (exclusion des outils installés
par défaut sur les ordinateurs ou déployés de façon standard par les TI) sur
les onglets de secteurs de
« User_Workstation_App_Associations_Consolidated.xlsx »
(jusqu'à l'onglet « Non défini » inclusivement).

Comportement (conforme à Regle_03_outils_installes_par_defaut.md) :
1. Génère d'abord la liste distincte des applications correspondant à la
   règle (aucun déplacement). Cette analyse est faite en lecture seule
   (read_only) pour être rapide et économe en mémoire.
2. Traite ensuite UN secteur à la fois : le classeur est rouvert pour cet
   onglet seulement, les lignes correspondantes sont déplacées vers l'onglet
   « Excluded Regle03 » (créé une seule fois) sans transformation des données
   d'origine (une colonne finale « Raison » est ajoutée dans l'onglet
   d'exclusion), puis le classeur est sauvegardé, fermé et les ressources
   sont libérées avant de passer au secteur suivant.
3. La progression est affichée à l'écran (onglet courant, temps écoulé et
   estimation du temps restant). Aucune intervention n'est requise entre
   les secteurs.

Correspondances visées (Règle 3) : 7-Zip, AI ExpertMeet (et Plugin),
FortiClient (VPN), FortiMonitor Agent, VLC, Adobe Acrobat Reader DC (version
gratuite), Antidote (incl. composantes Druide), ScreenConnect, FortiFone, VIA
(RoomDevice, VIAApp), XeroxCorp, Yubico.

Garde-fou (CONSERVÉES volontairement) :
- « Adobe Acrobat (64-bit) » : ligne représentative d'Acrobat présente dans
  les 18 secteurs. Les données ne permettant pas de distinguer l'Acrobat
  payant de la version gratuite, seule cette ligne est conservée en attente
  de confirmation (entrevues); « Adobe Acrobat » classique est exclu;
- « FortiToken » : non listé par le MD (seul FortiFone est ciblé);
- les composantes Adobe autres que celles du Reader gratuit (Photoshop,
  Illustrator, Creative Cloud, etc.).

Usage : py appliquer_regle03.py --dry-run   # simulation (aucune écriture)
        py appliquer_regle03.py             # exécution réelle
"""

import argparse
import gc
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from colonnes import aligner, index_colonnes, valeur

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
EXCLUDED_SHEET = "Excluded Regle03"
REASON = "Règle 3 — Outil installé par défaut / déployé standard par les TI"
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
IGNORED = {"logiciels", "sommaire", "feuil1", "excluded regle01", "excluded regle02",
           EXCLUDED_SHEET.casefold()}

# Légende des catégories affichées dans la liste distincte.
CATEGORIES_LEGENDE = ("A=7-Zip, B=AI ExpertMeet, C=FortiClient(VPN)/Monitor, "
                      "D=VLC, E=FortiFone, F=ScreenConnect, "
                      "G=Antidote/Druide, H=VIA, I=Xerox, J=Yubico, "
                      "K=Acrobat Reader (gratuit)")


def _norm(value):
    """Normalise une valeur : casefold, apostrophes unifiées, espaces réduits."""
    text = "" if value is None else str(value)
    for ch in ("’", "‘", "`"):
        text = text.replace(ch, "'")
    return " ".join(text.split()).casefold()


# --- Garde-fou Acrobat : ligne représentative CONSERVÉE (vérifiée avant tout)
# « Adobe Acrobat (64-bit) » est la ligne représentative d'Acrobat présente
# dans les 18 secteurs; ne pas la confondre avec « adobe acrobat » (exact),
# qui est exclu par la catégorie K.
_KEEP_ACROBAT_EXACT = {"adobe acrobat (64-bit)"}

# --- A) 7-Zip ----------------------------------------------------------------
_ZIP7_SUBSTRINGS = (
    "7-zip",                 # 7-Zip 18.05 / 25.01 / 26.x (x64), pub. Igor Pavlov
)

# --- B) AI ExpertMeet (et AI ExpertMeet Plugin) -------------------------------
_EXPERTMEET_SUBSTRINGS = (
    "expertmeet",            # AI ExpertMeet x.y.z / AI ExpertMeet Plugin (ASUSTeK)
)

# --- C) FortiClient / FortiClient VPN / FortiMonitor Agent ----------------------
# « fortimonitor » a été ajouté le 2026-02-10 : ces motifs figuraient
# auparavant dans la Règle 1, d'où ils ont été retirés de _EXTRAS_SUBSTRINGS
# dans appliquer_regle01.py. Ces motifs ont été déplacés vers cette règle, qui
# décrit les outils déployés par les TI (ce qu'est FortiMonitor Agent) et non
# les composantes livrées avec Windows 11. Sans cet ajout, les lignes
# « FortiMonitor Agent » ne seraient exclues par AUCUNE règle et
# réapparaîtraient dans les onglets de secteurs.
# « FortiToken » reste volontairement NON ciblé (garde-fou, cf. docstring).
_FORTI_SUBSTRINGS = (
    "forticlient",           # FortiClient, FortiClient VPN
    "fortimonitor",          # FortiMonitor Agent (ajout validé 2026-02-10)
)

# --- D) VLC --------------------------------------------------------------------
_VLC_RE = re.compile(r"\bvlc\b")   # VLC media player (VideoLAN)

# --- E) FortiFone (FortiToken volontairement CONSERVÉ : non listé au MD) ------
_FORTIFONE_SUBSTRINGS = (
    "fortifone",
)

# --- F) ScreenConnect (nom OU publisher) --------------------------------------
# « Connectwize » et « ConnectWise ScreenConnect »,
# pub. wendake.screenconnect.com : le nom seul ne suffit pas pour Connectwize.
_SCREENCONNECT_SUBSTRINGS = (
    "screenconnect",
)

# --- G) Antidote (incl. composantes Druide) ------------------------------------
_ANTIDOTE_SUBSTRINGS = (
    "antidote",              # Antidote 11/12, English module
)
_DRUIDE_PUB_SUBSTRINGS = (
    "druide",                # capte « Assistant de téléchargement » (Druide inc.)
)

# --- H) VIA (RoomDevice, VIAApp) -----------------------------------------------
# Début du nom : VIAApp, VIADisplayDriver (Kramer), VIA RoomDevice.
_VIA_RE = re.compile(r"^via")

# --- I) XeroxCorp (nom OU publisher) -------------------------------------------
_XEROX_SUBSTRINGS = (
    "xerox",                 # Xerox Desktop / Print and Scan Experience, Pilote Xerox W100
)

# --- J) Yubico ------------------------------------------------------------------
_YUBICO_SUBSTRINGS = (
    "yubico",                # Yubico Authenticator
)

# --- K) Adobe Acrobat Reader DC (version gratuite) et composantes Reader ------
# « Adobe Acrobat (64-bit) » est conservée (garde-fou); « Adobe Acrobat »
# classique est exclu (une seule ligne représentative est conservée, cf. MD).
_ACROBAT_EXACT = {
    "adobe acrobat",
}
_ACROBAT_READER_SUBSTRINGS = (
    "acrobat reader",              # Adobe Acrobat Reader, - Français,
                                   # Asian Language Dictionaries Support For Adobe Acrobat Reader
    "reader notification client",  # composante du Reader gratuit (pub. Adobe)
)

def match_rule3(logiciel, manufacturier):
    """Retourne la catégorie ('A' à 'K') si la ligne correspond à la
    Règle 3, sinon None."""
    name = _norm(logiciel)
    if not name:
        return None
    pub = _norm(manufacturier)

    # Garde-fou : « Adobe Acrobat (64-bit) » (ligne représentative présente
    # sur tous les postes) est toujours conservée.
    if name in _KEEP_ACROBAT_EXACT:
        return None

    if any(s in name for s in _ZIP7_SUBSTRINGS):
        return "A"
    if any(s in name for s in _EXPERTMEET_SUBSTRINGS):
        return "B"
    if any(s in name for s in _FORTI_SUBSTRINGS):
        return "C"
    if _VLC_RE.search(name):
        return "D"
    if any(s in name for s in _FORTIFONE_SUBSTRINGS):
        return "E"
    if (any(s in name for s in _SCREENCONNECT_SUBSTRINGS)
            or any(s in pub for s in _SCREENCONNECT_SUBSTRINGS)):
        return "F"
    if (any(s in name for s in _ANTIDOTE_SUBSTRINGS)
            or any(s in pub for s in _DRUIDE_PUB_SUBSTRINGS)):
        return "G"
    if _VIA_RE.match(name):
        return "H"
    if (any(s in name for s in _XEROX_SUBSTRINGS)
            or any(s in pub for s in _XEROX_SUBSTRINGS)):
        return "I"
    if any(s in name for s in _YUBICO_SUBSTRINGS):
        return "J"
    if name in _ACROBAT_EXACT or any(s in name for s in _ACROBAT_READER_SUBSTRINGS):
        return "K"
    return None


def _secteurs(sheetnames):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement."""
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in sheetnames:
        if name.casefold() not in ignored:
            secteurs.append(name)
        if name.casefold() == LAST_SECTOR.casefold():
            break
    return secteurs


def _fmt_duree(secondes):
    """Formate une durée en secondes sous la forme « 1 h 02 min 03 s »."""
    secondes = int(secondes)
    h, reste = divmod(secondes, 3600)
    m, s = divmod(reste, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"

def main():
    # Évite les UnicodeEncodeError quand la sortie est redirigée vers un
    # fichier (encodage console cp1252) : les caractères non représentables
    # sont remplacés plutôt que de faire planter le script.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(
        description="Applique la Règle 3 (outils installés par défaut).")
    parser.add_argument("--dry-run", action="store_true",
                        help="Simulation : affiche les résultats sans modifier le classeur.")
    args = parser.parse_args()

    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
        shutil.copy2(WORKBOOK, backup)
        print(f"Copie de sauvegarde : {backup.name}", flush=True)

    # Étape 1 — analyse en lecture seule (rapide, peu de mémoire) :
    # liste distincte des applications correspondant à la règle et nombre
    # de lignes à déplacer par secteur. Aucun déplacement à cette étape.
    print("\nÉtape 1/2 — Analyse des onglets (lecture seule)...", flush=True)
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    try:
        secteurs = _secteurs(wb.sheetnames)
        print(f"Onglets de secteurs à traiter (jusqu'à « {LAST_SECTOR} ») : "
              f"{len(secteurs)}", flush=True)
        header = None
        distinct = set()
        compteurs = {}   # {onglet: nb de lignes correspondantes}
        entetes = {}     # {onglet: en-tête d'origine}
        index_par_onglet = {}  # {onglet: index des colonnes requises}
        for sheet_name in secteurs:
            ws = wb[sheet_name]
            lignes = ws.iter_rows(values_only=True)
            entete = next(lignes, None)
            if entete is not None:
                entetes[sheet_name] = list(entete)
                # Les colonnes sont localisées par leur LIBELLÉ, jamais par
                # leur position : l'ordre peut varier d'une extraction à
                # l'autre. Lève une RuntimeError explicite si une colonne
                # requise est absente.
                index_par_onglet[sheet_name] = index_colonnes(
                    entete, ("logiciel", "manufacturier"), onglet=sheet_name)
                if header is None:
                    header = list(entete)
            cols = index_par_onglet[sheet_name]
            count = 0
            for row in lignes:
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                categorie = match_rule3(logiciel, manufacturier)
                if categorie:
                    count += 1
                    distinct.add((str(logiciel).strip(), categorie))
            compteurs[sheet_name] = count
    finally:
        wb.close()
        del wb
        gc.collect()

    print(f"\nListe distincte des applications correspondant à la Règle 3 "
          f"({CATEGORIES_LEGENDE}) :")
    for logiciel, categorie in sorted(distinct, key=lambda t: (t[1], t[0].casefold())):
        print(f"  [{categorie}] {logiciel}")
    total_planned = sum(compteurs.values())
    print(f"\nApplications distinctes : {len(distinct)}")
    print(f"Lignes à déplacer (tous secteurs) : {total_planned}", flush=True)
    if args.dry_run:
        print("Mode --dry-run : aucune modification.")
        return

    # Étape 2 — un secteur à la fois : ouverture du classeur, déplacement des
    # lignes vers « Excluded Regle03 », sauvegarde, fermeture et libération
    # des ressources avant de passer au secteur suivant (aucune pause).
    print(f"\nÉtape 2/2 — Traitement des secteurs (un à la fois, "
          f"sauvegarde après chacun) :", flush=True)
    debut = time.perf_counter()
    total_deplace = 0
    nb_secteurs = len(secteurs)
    for i, sheet_name in enumerate(secteurs, start=1):
        ecoule = time.perf_counter() - debut
        if i > 1:
            restant = ecoule / (i - 1) * (nb_secteurs - i + 1)
            progression = (f"temps écoulé {_fmt_duree(ecoule)}, "
                           f"restant estimé {_fmt_duree(restant)}")
        else:
            progression = "estimation du temps restant après le 1er onglet"
        print(f"\n[{i}/{nb_secteurs}] Onglet « {sheet_name} » ({progression})",
              flush=True)

        if compteurs[sheet_name] == 0:
            print("  -> 0 ligne à déplacer, onglet inchangé.", flush=True)
            continue

        t0 = time.perf_counter()
        wb = openpyxl.load_workbook(WORKBOOK)
        try:
            ws = wb[sheet_name]
            cols = index_par_onglet[sheet_name]
            matches = []
            for row_idx, row in enumerate(
                    ws.iter_rows(min_row=2, values_only=True), start=2):
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                if match_rule3(logiciel, manufacturier):
                    matches.append((row_idx, row))

            excluded_ws = (wb[EXCLUDED_SHEET]
                           if EXCLUDED_SHEET in wb.sheetnames else None)
            if excluded_ws is None:
                excluded_ws = wb.create_sheet(EXCLUDED_SHEET)
                excluded_ws.append((header or []) + ["Raison"])
            for _row_idx, values in matches:
                # Réordonne la ligne vers l'en-tête de référence : les onglets
                # de secteurs peuvent présenter un ordre de colonnes différent
                # sans que l'onglet d'exclusion ne scramblie les colonnes.
                excluded_ws.append(
                    aligner(values, entetes[sheet_name], header,
                            onglet=sheet_name) + [REASON])
            # Suppression par blocs de lignes consécutives (de bas en haut) :
            # beaucoup plus rapide qu'une suppression ligne par ligne.
            indices = sorted(m[0] for m in matches)
            i = len(indices) - 1
            while i >= 0:
                j = i
                while j > 0 and indices[j - 1] == indices[j] - 1:
                    j -= 1
                ws.delete_rows(indices[j], indices[i] - indices[j] + 1)
                i = j - 1

            # Ajuste les plages des tableaux structurés et des filtres à la
            # nouvelle fin de données. Sans cela, les plages pointent au-delà
            # des lignes restantes et Excel signale un contenu corrompu à
            # l'ouverture du fichier.
            fin = max(ws.max_row, 1)
            for tbl in ws.tables.values():
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", tbl.ref)
                if m and int(m.group(3)) != fin:
                    nouveau = f"{m.group(1)}:{m.group(2)}{fin}"
                    tbl.ref = nouveau
                    if tbl.autoFilter is not None:
                        tbl.autoFilter.ref = nouveau
            if ws.auto_filter.ref:
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", ws.auto_filter.ref)
                if m and int(m.group(3)) != fin:
                    ws.auto_filter.ref = f"{m.group(1)}:{m.group(2)}{fin}"

            wb.save(WORKBOOK)
        finally:
            wb.close()
            del wb
            gc.collect()  # libère les ressources avant le secteur suivant

        total_deplace += len(matches)
        print(f"  -> {len(matches)} ligne(s) déplacée(s) et sauvegardée(s) "
              f"en {_fmt_duree(time.perf_counter() - t0)}.", flush=True)

    print(f"\nTerminé en {_fmt_duree(time.perf_counter() - debut)}.")
    print(f"Total déplacé : {total_deplace} ligne(s) -> « {EXCLUDED_SHEET} ».")
    print("Aucune perte de données : chaque ligne retirée y est copiée telle quelle.")


if __name__ == "__main__":
    main()
