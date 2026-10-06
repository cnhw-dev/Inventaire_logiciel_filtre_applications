"""
top10_applications.py — Calcule le « top 10 » des applications les plus
fréquentes selon le « Nom d'hôte » et l'écrit dans un onglet « Top 10 » de
« User_Workstation_App_Associations_Consolidated.xlsx ».

Méthodologie :
- les onglets de secteurs (jusqu'à « Non défini » inclusivement) sont
  regroupés; les onglets « Feuil1 », « Logiciels », « Sommaire »,
  « Excluded RegleXX » et « Top 10 » sont ignorés;
- pour chaque logiciel, on compte le nombre de postes (« Ordinateur »)
  DISTINCTS (une application installée deux fois sur le même poste ne
  compte qu'une fois);
- le manufacturier affiché est la valeur la plus fréquente (mode) pour le
  logiciel;
- le pourcentage est calculé sur le nombre total de postes distincts.

L'onglet « Top 10 » est inséré immédiatement après l'onglet « Sommaire »
(ou à la fin du classeur si celui-ci est absent). S'il existe déjà, il
est remplacé.

Comportement (calqué sur appliquer_regle05.py) :
1. Une copie de sauvegarde horodatée du classeur est créée avant toute
   écriture (sauf en simulation).
2. La progression est affichée à l'écran. Aucune intervention n'est requise.

Usage : py top10_applications.py --dry-run   # simulation (aucune écriture)
        py top10_applications.py             # exécution réelle
"""

import argparse
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
TOP10_SHEET = "Top 10"
AFTER_SHEET = "Sommaire"  # l'onglet « Top 10 » est inséré après celui-ci
IGNORED = {"logiciels", "sommaire", "feuil1", TOP10_SHEET.casefold()}

COL_LOGICIEL = "Logiciel"
COL_MANUFACTURIER = "Manufacturier"
COL_HOTE = "Ordinateur"

HEADERS = ["Rang", "Logiciel", "Manufacturier", "Nb d'hôtes", "% des hôtes"]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", start_color="4472C4")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


def secteurs(workbook_path):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement
    (Feuil1, Sommaire, Excluded RegleXX et Top 10 ignorés)."""
    noms = pd.ExcelFile(workbook_path).sheet_names
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in noms:
        nom = name.casefold()
        if nom not in ignored and not nom.startswith("excluded"):
            secteurs.append(name)
        if nom == LAST_SECTOR.casefold():
            break
    return secteurs


def calculer_top10(workbook_path, n=10):
    """Calcule le top N des logiciels par nombre d'hôtes distincts."""
    onglets = secteurs(workbook_path)
    df = pd.concat(
        [pd.read_excel(workbook_path, sheet_name=s) for s in onglets],
        ignore_index=True,
    )
    df = df.dropna(subset=[COL_LOGICIEL, COL_HOTE])

    total_hotes = df[COL_HOTE].nunique()
    stats = df.groupby(COL_LOGICIEL).agg(
        hotes=(COL_HOTE, "nunique"),
        manufacturier=(
            COL_MANUFACTURIER,
            lambda x: x.mode().iat[0] if len(x.mode()) else "",
        ),
    )
    stats = stats.sort_values("hotes", ascending=False).head(n).reset_index()

    lignes = []
    for rang, (_, row) in enumerate(stats.iterrows(), start=1):
        lignes.append(
            [
                rang,
                row[COL_LOGICIEL],
                row["manufacturier"],
                int(row["hotes"]),
                row["hotes"] / total_hotes if total_hotes else 0.0,
            ]
        )
    return lignes, total_hotes, len(onglets), len(df)


def ecrire_onglet(workbook_path, lignes):
    """Écrit (ou remplace) l'onglet « Top 10 » dans le classeur."""
    wb = openpyxl.load_workbook(workbook_path)
    if TOP10_SHEET in wb.sheetnames:
        del wb[TOP10_SHEET]

    if AFTER_SHEET in wb.sheetnames:
        index = wb.sheetnames.index(AFTER_SHEET) + 1
    else:
        index = len(wb.sheetnames)
    ws = wb.create_sheet(TOP10_SHEET, index)

    # En-tête
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    # Données
    for ligne in lignes:
        ws.append(ligne)

    # Mise en forme
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.0 %"
    largeurs = {1: 8, 2: 40, 3: 28, 4: 14, 5: 14}
    for col, largeur in largeurs.items():
        ws.column_dimensions[get_column_letter(col)].width = largeur
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    wb.save(workbook_path)
    wb.close()


def main():
    parser = argparse.ArgumentParser(
        description="Ajoute un onglet « Top 10 » des applications les plus "
        "fréquentes (par « Nom d'hôte » distinct) au classeur consolidé."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Simulation : affiche le top 10 sans modifier le classeur.",
    )
    args = parser.parse_args()

    print(f"Classeur : {WORKBOOK.name}", flush=True)
    lignes, total_hotes, nb_onglets, nb_lignes = calculer_top10(WORKBOOK)

    print(
        f"Données analysées : {nb_onglets} onglets de secteurs, "
        f"{nb_lignes} lignes, {total_hotes} hôtes distincts.",
        flush=True,
    )
    print(flush=True)
    entete = f"{'Rang':<5}{'Logiciel':<42}{'Manufacturier':<28}{'Hôtes':>7}{'%':>8}"
    print(entete, flush=True)
    print("-" * 90, flush=True)
    for rang, logiciel, manufacturier, hotes, pct in lignes:
        print(
            f"{rang:<5}{logiciel:<42}{manufacturier:<28}{hotes:>7}{pct:>8.1%}",
            flush=True,
        )

    if args.dry_run:
        print(flush=True)
        print("Simulation (--dry-run) : aucune écriture effectuée.", flush=True)
        return

    # Copie de sauvegarde horodatée avant écriture
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
    shutil.copy2(WORKBOOK, backup)
    print(flush=True)
    print(f"Copie de sauvegarde : {backup.name}", flush=True)

    ecrire_onglet(WORKBOOK, lignes)
    print(f"Onglet « {TOP10_SHEET} » écrit dans {WORKBOOK.name}.", flush=True)


if __name__ == "__main__":
    main()

