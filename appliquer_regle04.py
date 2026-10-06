"""
appliquer_regle04.py — Applique la Règle 4 (exclusion des pilotes,
utilitaires matériels et composantes système) sur les onglets de secteurs de
« User_Workstation_App_Associations_Consolidated.xlsx »
(jusqu'à l'onglet « Non défini » inclusivement).

Comportement (conforme à Regle_04_pilotes_utilitaires_materiels.md) :
1. Génère d'abord la liste distincte des applications correspondant à la
   règle (aucun déplacement). Cette analyse est faite en lecture seule
   (read_only) pour être rapide et économe en mémoire.
2. Traite ensuite UN secteur à la fois : le classeur est rouvert pour cet
   onglet seulement, les lignes correspondantes sont déplacées vers l'onglet
   « Excluded Regle04 » (créé une seule fois) sans transformation des données
   d'origine (une colonne finale « Raison » est ajoutée dans l'onglet
   d'exclusion), puis le classeur est sauvegardé, fermé et les ressources
   sont libérées avant de passer au secteur suivant.
3. La progression est affichée à l'écran (onglet courant, temps écoulé et
   estimation du temps restant). Aucune intervention n'est requise entre
   les secteurs.

Catégories visées (Règle 4) :
A=Composantes audio (Realtek, Dolby, Nahimic, C-Media, Dirac, Elevoc...),
B=Clavier/souris/pointage (Logitech, Kensington, ELAN, Synaptics, Razer...),
C=Pilotes génériques (nom contenant « driver »/« pilote », Windows Driver
  Package, Package de pilotes Windows, minidriver...),
D=Lenovo, E=ASUS, F=HP (incl. HP Wolf Security et HP Sure *), G=MSI,
H=AMD, I=Intel, J=NVIDIA,
K=Imprimantes/scanners et périphériques d'impression (Brother, Canon,
  PFU/Fujitsu, Zebra, DYMO, Evolis, Visioneer, MakerBot, imprimantes PDF
  virtuelles Bullzip/CutePDF/Ghostscript...),
L=Autres périphériques/matériel (Garmin, Apple Mobile Device Support,
  Bonjour, DJI, Wacom, X-Rite, utilitaires RGB/SSD/micrologiciel...),
M=Services système Windows détectés comme applications (« * Session »,
  SessionManager, Narrator Extension, EpmShellExtension, outils de
  déploiement Windows, Sécurité Windows, OS « Microsoft Windows 10/11 »...),
N=Runtimes/bibliothèques système (MSVCRT, D3DX10, .NET SDK/Packs/Workloads,
  Visual Studio Tools for Office, Silverlight, Java Update...),
O=Lignes au nom GUID ({...}).

Garde-fous (CONSERVÉES volontairement) :
- « Poly Lens », « Poly Studio » et « Poly Lens Control Service » (publisher
  HP Inc.), « Jabra Direct », « kuando Busylight/HUB » : périphériques audio
  conservés (décision de portée);
- la connectivité BD métier : Microsoft ODBC/OLE DB Driver pour SQL Server,
  MySQL Connector/ODBC, SQL Server Native Client/Compact/LocalDB, Crystal
  Reports / CR115_rdc_runtime (le mot « Driver » ou « .NET » ne doit pas
  les exclure — requis par les applications métier);
- « HYP2003-RAMQ » (middleware jeton RAMQ lié aux applications DSQ),
  McAfee/WebAdvisor (sécurité, hors portée);
- les outils de développement : Python, Node.js, Docker, VS Code,
  IIS Express, WSL, Dev Home;
- les applications métier/utilisateur liées à du matériel mais applicatives :
  IPEVO Visualizer, AmScope, iVMS-4200/SADP (Hikvision), CDVI Module Finder,
  Cirrus SMS/NoiseTools, Analook/Kaleidoscope, ArcSoft Print Creations,
  iTunes (seuls ses services Apple sont exclus);
- les composantes Citrix, Adobe, Autodesk et SMART Notebook (seul
  « SMART Product Drivers » est exclu via la catégorie C).

Usage : py appliquer_regle04.py --dry-run   # simulation (aucune écriture)
        py appliquer_regle04.py             # exécution réelle
"""

import argparse
import gc
import re
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

from colonnes import aligner, index_colonnes, valeur

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK = BASE_DIR / "User_Workstation_App_Associations_Consolidated.xlsx"
EXCLUDED_SHEET = "Excluded Regle04"
REASON = "Règle 4 — Pilote / utilitaire matériel / composante système"
LAST_SECTOR = "Non défini"  # dernier onglet de secteur à traiter (inclusivement)
IGNORED = {"logiciels", "sommaire", "feuil1", "excluded regle01", "excluded regle02",
           "excluded regle03", EXCLUDED_SHEET.casefold()}

# Légende des catégories affichées dans la liste distincte.
CATEGORIES_LEGENDE = ("A=Audio, B=Clavier/souris, C=Pilotes, D=Lenovo, "
                      "E=ASUS, F=HP, G=MSI, H=AMD, I=Intel, J=NVIDIA, "
                      "K=Imprimantes/scanners, L=Périphériques divers, "
                      "M=Services système Windows, N=Runtimes système, "
                      "O=Noms GUID")


def _norm(value):
    """Normalise une valeur : casefold, apostrophes unifiées, espaces réduits."""
    text = "" if value is None else str(value)
    for ch in ("’", "‘", "`"):
        text = text.replace(ch, "'")
    return " ".join(text.split()).casefold()


# --- Garde-fous (vérifiés avant toute catégorie) ----------------------------
# Poly (publisher HP Inc.) : périphériques audio conservés (décision de portée).
# Connectivité BD métier : ne pas exclure malgré « Driver »/« .NET » dans le nom.
_KEEP_PREFIXES = ("poly ", "cr115")
_KEEP_SUBSTRINGS = ("odbc", "ole db", "sql server", "mysql", "crystal reports")

# --- A) Composantes audio -----------------------------------------------------
_AUDIO_NAME = ("realtek", "dolby", "nahimic", "audiowizard", "c-media",
               "dirac audio", "elevoc", "smartaudio",
               "smart microphone settings", "hp audio control", "a-volute")
_AUDIO_PUB = ("realtek", "a-volute", "icepower")

# --- B) Clavier / souris / pointage -------------------------------------------
_POINTAGE_NAME = ("logitech", "kensingtonworks", "elan ", "trackpoint",
                  "synaptics", "razer", "xbox 360 controller emulator",
                  "hp programmable key", "touchpad")
_POINTAGE_PUB = ("logitech", "elan microelectronics", "synaptics", "razer",
                 "kensington")
_POINTAGE_PREFIX = ("logi ",)   # Logi Options+, Logi Plugin Service

# --- C) Pilotes génériques ------------------------------------------------------
# « Windows Driver Package - ... », « Package de pilotes Windows - ... »,
# Brother/Canon/Garmin/Munbyn Drivers, NVIDIA Pilote graphique, etc.
_PILOTE_RE = re.compile(r"\b(drivers?|pilotes?)\b")
# Cas sans espace autour du mot (pas de frontière \b) : SafeNet Minidriver,
# DFUDriverSetupX64Setup, ISS_Drivers_x64, AMD_Chipset_Drivers.
_PILOTE_EXTRA = ("minidriver", "dfudriver", "iss_drivers", "amd_chipset_drivers")

# --- D) Lenovo ------------------------------------------------------------------
_LENOVO_NAME = ("lenovo", "vantage", "glance by mirametrix", "memcmft",
                "ladmlauncherservice", "thinkpad")
_LENOVO_PUB = ("lenovo",)
# « Smart Note » (Lenovo) avec frontière de mot : ne pas capter
# « SMART Notebook » (application éducative conservée).
_LENOVO_SMART_NOTE_RE = re.compile(r"\bsmart note\b")

# --- E) ASUS --------------------------------------------------------------------
_ASUS_NAME = ("asus", "armoury crate", "glidex", "screenpad", "b9eced6f.",
              "gamesdk service")
_ASUS_AURA_RE = re.compile(r"\baura\b")   # AURA Service, Kingston AURA DRAM...
_ASUS_ROG_RE = re.compile(r"^rog\b")      # ROG Live Service, ROG OMNI RECEIVER
_ASUS_PUB = ("asustek",)

# --- F) HP (incl. HP Wolf Security, HP Sure *, plugins et services HP) ---------
_HP_RE = re.compile(r"^hp\b")             # HP Smart, HP Wolf Security, « HP »...
_HP_PREFIX = ("hpp", "hpclj")             # hppLaserJetService, HPCLJProMFPM177
_HP_PUB_EXACT = {"hp", "hp inc.", "hp inc"}
_HP_PUB_SUB = ("hewlett",)                # Hewlett-Packard, HPSSupply, MarketResearch

# --- G) MSI ---------------------------------------------------------------------
_MSI_RE = re.compile(r"^msi\b")           # MSI Center, MSI Driver & App Center
_MSI_PUB = ("micro-star", "msiappadm")
_MSI_PUB_EXACT = {"msi"}

# --- H) AMD ---------------------------------------------------------------------
# « amd » non suivi d'un caractère alphanumérique : AMD Software,
# AMD_Chipset_Drivers, « AMD » seul (capte aussi les versions Store dont
# le publisher est Microsoft).
_AMD_RE = re.compile(r"^amd(?![a-z0-9])")
_AMD_NAME = ("branding64", "ryzenmaster")
_AMD_PUB = ("advanced micro devices",)


# --- I) Intel -------------------------------------------------------------------
_INTEL_NAME = ("intel", "killer intelligence center", "appup.",
               "dynamic application loader host interface service")
_INTEL_PUB = ("intel",)

# --- J) NVIDIA ------------------------------------------------------------------
_NVIDIA_NAME = ("nvidia", "physx", "geforce")
_NVIDIA_RE = re.compile(r"^nv")           # NvCpl, NvModuleTracker
_NVIDIA_PUB = ("nvidia",)

# --- K) Imprimantes / scanners / périphériques d'impression ---------------------
_PRINTER_NAME = ("brother", "canon", "paperstream", "scansnap", "fi series",
                 "error recovery guide for fi", "software operation panel",
                 "scanner central admin", "network setup tool for fi",
                 "zebra", "evolis", "dymo", "cardpresso", "onetouch",
                 "visioneer", "samsung printer", "munbyn", "makerbot",
                 "desktopconverter", "controlcenter4", "statusmonitor",
                 "bullzip pdf printer", "cutepdf", "ghostscript",
                 "fujitsu", "iprint&scan")
_PRINTER_PUB = ("brother", "canon", "pfu", "fujitsu", "zebra", "dymo",
                "evolis", "visioneer", "makerbot", "stratasys",
                "newell rubbermaid")

# --- L) Autres périphériques / utilitaires matériels ----------------------------
_PERIPH_NAME = ("garmin", "apple mobile device support", "bonjour",
                "apple software update", "dji assistant", "lg monitor app",
                "x-rite", "mobile broadband", "deviceinstaller", "vli hub",
                "wd_black", "patriot viper", "universal holtek",
                "ene rgb hal", "ene_ehd", "phison", "hwmonitor",
                "ant drivers installer", "inreach sync")
_PERIPH_PUB = ("garmin", "huawei", "x-rite", "wacom", "ene tech",
               "phison", "patriot memory", "lantronix", "dynastream")

# --- M) Services système Windows détectés comme applications --------------------
_SYSWIN_SUFFIX = " session"               # Image Search Session, Text Recognition...
_SYSWIN_NAME = ("sessionmanager", "narrator extension", "narratorextension",
                "epmshellextension",
                "module d'expérience de fonctionnalités windows",
                "windows shell experience", "flux de connexion réseau",
                "bcd and boot", "oa3tool", "oacheck", "oatool", "oscdimg",
                "diagnosticshub", "pstokenizer", "windows deployment",
                "kit de déploiement et d'évaluation windows",
                "kits configuration installer", "windows system image manager",
                "windows ml runtime", "windows package manager source",
                "desktop app web viewer",
                "visionneuse web de l'appli d'appareil de bureau",
                "ppiprojection", "windowsholographicfirstrun",
                "defender on-device registry filter", "sécurité windows",
                "microsoftpcmanager", "photosai")
_SYSWIN_EXACT = {"uup"}
_SYSWIN_OS_RE = re.compile(r"^microsoft windows 1[01]\b")  # l'OS listé comme appli

# --- N) Runtimes / bibliothèques système ----------------------------------------
_RUNTIME_NAME = (".net", "msvcrt", "d3dx10", "visual studio 2010 tools",
                 "visual studio tools for applications",
                 "visual studio installer", "visual studio setup",
                 "silverlight", "microsoft advertising", "microsoft wse",
                 "java 8 update", "java auto updater",
                 "microsoft application error reporting",
                 "windows desktop targeting pack")
_RUNTIME_PREFIX = ("vs_",)                # vs_CoreEditorFonts, vs_community*...

# --- O) Noms GUID ---------------------------------------------------------------
_GUID_PREFIX = "{"


def match_rule4(logiciel, manufacturier):
    """Retourne la catégorie ('A' à 'O') si la ligne correspond à la
    Règle 4, sinon None."""
    name = _norm(logiciel)
    if not name:
        return None
    pub = _norm(manufacturier)

    # Garde-fous : périphériques audio conservés (Poly) et connectivité BD
    # métier (ODBC/OLE DB/SQL Server/MySQL/Crystal Reports).
    if name.startswith(_KEEP_PREFIXES):
        return None
    if any(s in name for s in _KEEP_SUBSTRINGS):
        return None

    # O) Lignes au nom GUID ({...}), tous publishers.
    if name.startswith(_GUID_PREFIX):
        return "O"

    if (any(s in name for s in _AUDIO_NAME)
            or any(s in pub for s in _AUDIO_PUB)):
        return "A"
    if (any(s in name for s in _POINTAGE_NAME)
            or name.startswith(_POINTAGE_PREFIX)
            or any(s in pub for s in _POINTAGE_PUB)):
        return "B"
    if _PILOTE_RE.search(name) or any(s in name for s in _PILOTE_EXTRA):
        return "C"
    if (any(s in name for s in _LENOVO_NAME)
            or _LENOVO_SMART_NOTE_RE.search(name)
            or any(s in pub for s in _LENOVO_PUB)):
        return "D"
    if (any(s in name for s in _ASUS_NAME)
            or _ASUS_AURA_RE.search(name) or _ASUS_ROG_RE.search(name)
            or any(s in pub for s in _ASUS_PUB)):
        return "E"
    if (_HP_RE.search(name) or name.startswith(_HP_PREFIX)
            or pub in _HP_PUB_EXACT
            or any(s in pub for s in _HP_PUB_SUB)):
        return "F"
    if (_MSI_RE.search(name) or pub in _MSI_PUB_EXACT
            or any(s in pub for s in _MSI_PUB)):
        return "G"
    if (_AMD_RE.search(name) or any(s in name for s in _AMD_NAME)
            or any(s in pub for s in _AMD_PUB)):
        return "H"
    if (any(s in name for s in _INTEL_NAME)
            or any(s in pub for s in _INTEL_PUB)):
        return "I"
    if (any(s in name for s in _NVIDIA_NAME) or _NVIDIA_RE.search(name)
            or any(s in pub for s in _NVIDIA_PUB)):
        return "J"
    if (any(s in name for s in _PRINTER_NAME)
            or any(s in pub for s in _PRINTER_PUB)):
        return "K"
    if (any(s in name for s in _PERIPH_NAME)
            or any(s in pub for s in _PERIPH_PUB)):
        return "L"
    if (name.endswith(_SYSWIN_SUFFIX) or name in _SYSWIN_EXACT
            or any(s in name for s in _SYSWIN_NAME)
            or _SYSWIN_OS_RE.search(name)):
        return "M"
    if (any(s in name for s in _RUNTIME_NAME)
            or name.startswith(_RUNTIME_PREFIX)):
        return "N"
    return None


def _secteurs(sheetnames):
    """Retourne les onglets de secteurs jusqu'à « Non défini » inclusivement."""
    ignored = {s.casefold() for s in IGNORED}
    secteurs = []
    for name in sheetnames:
        if name.casefold() not in ignored:
            secteurs.append(name)
        if name.casefold() == LAST_SECTOR.casefold():
            break
    return secteurs


def _fmt_duree(secondes):
    """Formate une durée en secondes sous la forme « 1 h 02 min 03 s »."""
    secondes = int(secondes)
    h, reste = divmod(secondes, 3600)
    m, s = divmod(reste, 60)
    if h:
        return f"{h} h {m:02d} min {s:02d} s"
    if m:
        return f"{m} min {s:02d} s"
    return f"{s} s"

def main():
    # Évite les UnicodeEncodeError quand la sortie est redirigée vers un
    # fichier (encodage console cp1252) : les caractères non représentables
    # sont remplacés plutôt que de faire planter le script.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(
        description="Applique la Règle 4 (pilotes, utilitaires matériels "
                    "et composantes système).")
    parser.add_argument("--dry-run", action="store_true",
                        help="Simulation : affiche les résultats sans modifier le classeur.")
    args = parser.parse_args()

    if not args.dry_run:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = WORKBOOK.with_name(f"{WORKBOOK.stem}_backup_{stamp}.xlsx")
        shutil.copy2(WORKBOOK, backup)
        print(f"Copie de sauvegarde : {backup.name}", flush=True)

    # Étape 1 — analyse en lecture seule (rapide, peu de mémoire) :
    # liste distincte des applications correspondant à la règle et nombre
    # de lignes à déplacer par secteur. Aucun déplacement à cette étape.
    print("\nÉtape 1/2 — Analyse des onglets (lecture seule)...", flush=True)
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    try:
        secteurs = _secteurs(wb.sheetnames)
        print(f"Onglets de secteurs à traiter (jusqu'à « {LAST_SECTOR} ») : "
              f"{len(secteurs)}", flush=True)
        header = None
        distinct = set()
        compteurs = {}   # {onglet: nb de lignes correspondantes}
        entetes = {}     # {onglet: en-tête d'origine}
        index_par_onglet = {}  # {onglet: index des colonnes requises}
        for sheet_name in secteurs:
            ws = wb[sheet_name]
            lignes = ws.iter_rows(values_only=True)
            entete = next(lignes, None)
            if entete is not None:
                entetes[sheet_name] = list(entete)
                # Les colonnes sont localisées par leur LIBELLÉ, jamais par
                # leur position : l'ordre peut varier d'une extraction à
                # l'autre. Lève une RuntimeError explicite si une colonne
                # requise est absente.
                index_par_onglet[sheet_name] = index_colonnes(
                    entete, ("logiciel", "manufacturier"), onglet=sheet_name)
                if header is None:
                    header = list(entete)
            cols = index_par_onglet[sheet_name]
            count = 0
            for row in lignes:
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                categorie = match_rule4(logiciel, manufacturier)
                if categorie:
                    count += 1
                    distinct.add((str(logiciel).strip(), categorie))
            compteurs[sheet_name] = count
    finally:
        wb.close()
        del wb
        gc.collect()

    print(f"\nListe distincte des applications correspondant à la Règle 4 "
          f"({CATEGORIES_LEGENDE}) :")
    for logiciel, categorie in sorted(distinct, key=lambda t: (t[1], t[0].casefold())):
        print(f"  [{categorie}] {logiciel}")
    total_planned = sum(compteurs.values())
    print(f"\nApplications distinctes : {len(distinct)}")
    print(f"Lignes à déplacer (tous secteurs) : {total_planned}", flush=True)
    if args.dry_run:
        print("Mode --dry-run : aucune modification.")
        return

    # Étape 2 — un secteur à la fois : ouverture du classeur, déplacement des
    # lignes vers « Excluded Regle04 », sauvegarde, fermeture et libération
    # des ressources avant de passer au secteur suivant (aucune pause).
    print(f"\nÉtape 2/2 — Traitement des secteurs (un à la fois, "
          f"sauvegarde après chacun) :", flush=True)
    debut = time.perf_counter()
    total_deplace = 0
    nb_secteurs = len(secteurs)
    for i, sheet_name in enumerate(secteurs, start=1):
        ecoule = time.perf_counter() - debut
        if i > 1:
            restant = ecoule / (i - 1) * (nb_secteurs - i + 1)
            progression = (f"temps écoulé {_fmt_duree(ecoule)}, "
                           f"restant estimé {_fmt_duree(restant)}")
        else:
            progression = "estimation du temps restant après le 1er onglet"
        print(f"\n[{i}/{nb_secteurs}] Onglet « {sheet_name} » ({progression})",
              flush=True)

        if compteurs[sheet_name] == 0:
            print("  -> 0 ligne à déplacer, onglet inchangé.", flush=True)
            continue

        t0 = time.perf_counter()
        wb = openpyxl.load_workbook(WORKBOOK)
        try:
            ws = wb[sheet_name]
            cols = index_par_onglet[sheet_name]
            matches = []
            for row_idx, row in enumerate(
                    ws.iter_rows(min_row=2, values_only=True), start=2):
                logiciel = valeur(row, cols["logiciel"])
                manufacturier = valeur(row, cols["manufacturier"])
                if match_rule4(logiciel, manufacturier):
                    matches.append((row_idx, row))

            excluded_ws = (wb[EXCLUDED_SHEET]
                           if EXCLUDED_SHEET in wb.sheetnames else None)
            if excluded_ws is None:
                excluded_ws = wb.create_sheet(EXCLUDED_SHEET)
                excluded_ws.append((header or []) + ["Raison"])
            for _row_idx, values in matches:
                # Réordonne la ligne vers l'en-tête de référence : les onglets
                # de secteurs peuvent présenter un ordre de colonnes différent
                # sans que l'onglet d'exclusion ne scramblie les colonnes.
                excluded_ws.append(
                    aligner(values, entetes[sheet_name], header,
                            onglet=sheet_name) + [REASON])
            # Suppression par blocs de lignes consécutives (de bas en haut) :
            # beaucoup plus rapide qu'une suppression ligne par ligne.
            indices = sorted(m[0] for m in matches)
            i = len(indices) - 1
            while i >= 0:
                j = i
                while j > 0 and indices[j - 1] == indices[j] - 1:
                    j -= 1
                ws.delete_rows(indices[j], indices[i] - indices[j] + 1)
                i = j - 1

            # Ajuste les plages des tableaux structurés et des filtres à la
            # nouvelle fin de données. Sans cela, les plages pointent au-delà
            # des lignes restantes et Excel signale un contenu corrompu à
            # l'ouverture du fichier.
            fin = max(ws.max_row, 1)
            for tbl in ws.tables.values():
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", tbl.ref)
                if m and int(m.group(3)) != fin:
                    nouveau = f"{m.group(1)}:{m.group(2)}{fin}"
                    tbl.ref = nouveau
                    if tbl.autoFilter is not None:
                        tbl.autoFilter.ref = nouveau
            if ws.auto_filter.ref:
                m = re.match(r"^([A-Z]+\d+):([A-Z]+)(\d+)$", ws.auto_filter.ref)
                if m and int(m.group(3)) != fin:
                    ws.auto_filter.ref = f"{m.group(1)}:{m.group(2)}{fin}"

            wb.save(WORKBOOK)
        finally:
            wb.close()
            del wb
            gc.collect()  # libère les ressources avant le secteur suivant

        total_deplace += len(matches)
        print(f"  -> {len(matches)} ligne(s) déplacée(s) et sauvegardée(s) "
              f"en {_fmt_duree(time.perf_counter() - t0)}.", flush=True)

    print(f"\nTerminé en {_fmt_duree(time.perf_counter() - debut)}.")
    print(f"Total déplacé : {total_deplace} ligne(s) -> « {EXCLUDED_SHEET} ».")
    print("Aucune perte de données : chaque ligne retirée y est copiée telle quelle.")


if __name__ == "__main__":
    main()


